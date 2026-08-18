import io
import tempfile
import unittest
from pathlib import Path

from aurora_monitor.db import Database
from aurora_monitor.document_parser import ParsedDocument, parse_document
from aurora_monitor.fetcher import FetchResult
from aurora_monitor.monitor import Monitor, _looks_like_attachment
from aurora_monitor.positions import (
    UNKNOWN,
    _looks_like_major,
    _rows_from_text,
    extract_positions,
    upsert_positions,
)


def build_position_xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in [
        ["测试单位2026年公开招聘岗位简介表", "", "", "", "", "", "", "", "", ""],
        ["序号", "单位", "岗位\n代码", "岗位简介", "招聘\n人数", "开考比例", "岗位要求", "", "", "联系电话"],
        ["", "", "", "", "", "", "学历", "专业", "其他条件", ""],
        ["01", "测试中心", "01", "从事管理工作", "1", "1:3", "本科及以上", "法律类",
         "1.取得相应学位；2.中共（预备）党员；3.具有2年及以上基层工作经历；4.适合男性。", "123"],
        ["", "", "02", "从事财务工作", "2", "1:3", "本科及以上", "财务财会类", "1.取得相应学位。", ""],
    ]:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class PositionExtractionTests(unittest.TestCase):
    def test_xlsx_two_row_header_with_continuation(self):
        body = build_position_xlsx()
        parsed = parse_document(body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "https://example.gov.cn/table.xlsx")
        records = extract_positions(parsed, "https://example.gov.cn/table.xlsx")
        self.assertEqual(len(records), 2)
        first, second = records
        self.assertEqual(first.employer, "测试中心")
        self.assertEqual(first.position_code, "01")
        self.assertEqual(first.position_name, "从事管理工作")
        self.assertEqual(first.headcount, "1")
        self.assertEqual(first.education, "本科及以上")
        self.assertEqual(first.major_requirement, "法律类")
        self.assertEqual(first.degree, "取得相应学位")
        self.assertEqual(first.political_requirement, "中共（预备）党员")
        self.assertEqual(first.grassroots_requirement, "具有2年及以上基层工作经历")
        self.assertEqual(first.gender_requirement, "适合男性")
        self.assertEqual(second.employer, "测试中心")
        self.assertEqual(second.position_code, "02")
        self.assertEqual(second.headcount, "2")
        self.assertEqual(second.gender_requirement, UNKNOWN)

    def test_process_list_with_names_is_not_a_position_table(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["序号", "单位名称", "岗位名称", "招聘人数", "姓名", "总成绩", "排名"])
        sheet.append(["1", "测试局", "工作人员", "1", "张三", "70.1", "1"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        parsed = parse_document(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "https://example.gov.cn/list.xlsx")
        self.assertEqual(extract_positions(parsed, "https://example.gov.cn/list.xlsx"), [])

    def test_pdf_text_rows_are_extracted(self):
        text = (
            "某单位2026年招聘职位表\n"
            "职位代码\t单位名称\t职位名称\t学历要求\t专业要求\t招聘人数\n"
            "001\t测试局\t综合管理\t本科及以上\t法学\t1\n"
            "002\t测试局\t信息技术\t硕士研究生\t计算机科学与技术\t2\n"
        )
        parsed = ParsedDocument(text=text)
        records = extract_positions(parsed, "https://example.gov.cn/table.pdf")
        self.assertEqual([record.position_code for record in records], ["001", "002"])
        self.assertEqual(records[0].employer, "测试局")
        self.assertEqual(records[0].education, "本科及以上")
        self.assertEqual(records[1].major_requirement, "计算机科学与技术")

    def test_html_text_is_not_parsed_as_positions(self):
        parsed = ParsedDocument(text="职位代码 单位名称 职位名称 学历要求\n001 测试局 综合管理 本科")
        self.assertEqual(extract_positions(parsed, "https://example.gov.cn/notice.html"), [])

    def test_major_noise_is_rejected(self):
        self.assertTrue(_looks_like_major("电力相关专业"))
        self.assertTrue(_looks_like_major("法学专业"))
        self.assertFalse(_looks_like_major("在处理劳动争议时能保持冷静、专业"))
        self.assertFalse(_looks_like_major("能快速掌握相关专业知识与技能"))
        self.assertFalse(_looks_like_major("药学、食品、机械电子等相关专业优先"))

    def test_rows_from_text_requires_multiple_cells(self):
        rows = _rows_from_text("标题行\n职位代码\t单位\t人数\n001\t测试局\t1")
        self.assertEqual(len(rows), 2)

    def test_upsert_positions_is_idempotent(self):
        with tempfile.TemporaryDirectory() as directory:
            db = Database(Path(directory) / "a.db")
            db.init_schema()
            db.seed_regions()
            db.upsert_policy("p", 1, ["招聘"], [], [])
            db.add_source({
                "id": "s",
                "source_group": "jiangsu_city_recruitment",
                "region_code": "JS-南京",
                "publisher": "测试来源",
                "entry_url": "https://example.gov.cn/list",
                "keyword_policy_id": "p",
            })
            with db.transaction() as conn:
                conn.execute(
                    """INSERT INTO notice(id, source_id, title, normalized_title, url, decision, matched_terms, detail_status)
                       VALUES ('n1', 's', '招聘公告', '招聘', 'https://example.gov.cn/n1', 'candidate', '[]', 'fetched')"""
                )
                conn.execute(
                    """INSERT INTO evidence_version(id, notice_id, source_url, content_sha256)
                       VALUES ('e1', 'n1', 'https://example.gov.cn/a.xlsx', ?)""",
                    ("a" * 64,),
                )
            parsed = parse_document(build_position_xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "https://example.gov.cn/a.xlsx")
            records = extract_positions(parsed, "https://example.gov.cn/a.xlsx")
            self.assertEqual(upsert_positions(db, "n1", "e1", records), 2)
            self.assertEqual(upsert_positions(db, "n1", "e1", records), 2)
            self.assertEqual(db.connection.execute("SELECT COUNT(*) FROM position").fetchone()[0], 2)
            db.close()

    def test_monitor_stores_positions_with_evidence(self):
        with tempfile.TemporaryDirectory() as directory:
            db = Database(Path(directory) / "a.db")
            db.init_schema()
            db.seed_regions()
            db.upsert_policy("p", 1, ["招聘"], [], [])
            db.add_source({
                "id": "s",
                "source_group": "jiangsu_city_recruitment",
                "region_code": "JS-南京",
                "publisher": "测试来源",
                "entry_url": "https://example.gov.cn/list",
                "keyword_policy_id": "p",
            })
            with db.transaction() as conn:
                conn.execute(
                    """INSERT INTO notice(id, source_id, title, normalized_title, url, decision, matched_terms, detail_status)
                       VALUES ('n1', 's', '招聘公告', '招聘', 'https://example.gov.cn/n1', 'candidate', '[]', 'pending')"""
                )
            result = FetchResult(
                "https://example.gov.cn/a.xlsx", 200,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                build_position_xlsx(), 1,
            )
            _, inserted = Monitor(db)._store_evidence("n1", result)
            self.assertTrue(inserted)
            self.assertEqual(db.connection.execute("SELECT COUNT(*) FROM position").fetchone()[0], 2)
            self.assertEqual(db.connection.execute("SELECT detail_status FROM notice WHERE id='n1'").fetchone()[0], "fetched")
            db.close()

    def test_attachment_detection_covers_gateway_download_links(self):
        gateway_url = (
            "https://hrss.yangzhou.gov.cn/api-gateway/jpaas-web-server/front/document/download"
            "?fileUrl=abc%3D&fileName=2026%E5%B9%B4%E5%B2%97%E4%BD%8D%E6%9D%A1%E4%BB%B6%E7%AE%80%E4%BB%8B%E8%A1%A8.xls"
        )
        self.assertTrue(_looks_like_attachment("2026年扬州市市属事业单位统一公开招聘工作人员岗位条件简介表.xls", gateway_url))
        self.assertTrue(_looks_like_attachment("", "https://example.gov.cn/files/a.xlsx"))
        self.assertTrue(_looks_like_attachment("招聘岗位简介表", "https://example.gov.cn/view.html"))
        self.assertFalse(_looks_like_attachment("招聘公告", "https://example.gov.cn/view.html"))
        self.assertFalse(_looks_like_attachment("公告", "https://example.gov.cn/download?fileId=123"))

    def test_corrupt_xls_does_not_crash_parser(self):
        parsed = parse_document(b"not a real xls", "application/vnd.ms-excel", "https://example.gov.cn/a.xls")
        self.assertIn(parsed.parser_status, {"error", "needs_dependency"})
        self.assertTrue(parsed.warnings)


if __name__ == "__main__":
    unittest.main()
